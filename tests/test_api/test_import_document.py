import io
import os

import pytest
from docx import Document
from docx.shared import Inches

from src.config import get_settings
from src.db.engine import init_db, reset_db_state
from tests.test_api.auth_helpers import register_user_via_verification
from tests.test_document_import.builders import PdfText, build_pdf, docx_bytes, image_bytes

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def import_db(monkeypatch):
    db_path = os.path.join(os.path.dirname(__file__), "_test_import_document.db")
    if os.path.exists(db_path):
        os.remove(db_path)
    url = f"sqlite:///{db_path}"
    monkeypatch.setenv("DATABASE_URL", url)
    monkeypatch.setenv("DIRECT_DATABASE_URL", url)
    monkeypatch.setenv("AUTH_SECRET_KEY", "test-secret-key-for-jwt-signing-32chars")
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.setenv("SMTP_HOST", "")
    monkeypatch.setenv("SMTP_FROM", "")
    monkeypatch.setenv("SMTP_USER", "")
    monkeypatch.setenv("SMTP_PASSWORD", "")
    get_settings.cache_clear()
    reset_db_state()
    init_db()
    yield
    reset_db_state()
    get_settings.cache_clear()
    if os.path.exists(db_path):
        try:
            os.remove(db_path)
        except OSError:
            pass


async def _auth_headers(client, email: str = "importer@uni.edu") -> dict[str, str]:
    body = await register_user_via_verification(client, name="Importer", email=email, password="SecurePass1")
    return {"Authorization": f"Bearer {body['access_token']}"}


def _sample_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Imported Paper", style="Title")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Body with 10% & more.")
    doc.add_picture(io.BytesIO(image_bytes("PNG")), width=Inches(2))
    return docx_bytes(doc)


async def _upload(client, headers, filename: str, data: bytes, mime: str):
    return await client.post(
        "/api/v1/import/document",
        headers=headers,
        files={"file": (filename, data, mime)},
    )


@pytest.mark.asyncio
async def test_import_requires_authentication(client, import_db):
    response = await _upload(client, {}, "paper.docx", _sample_docx(), DOCX_MIME)
    assert response.status_code == 401


@pytest.mark.asyncio
async def test_import_docx_returns_project_payload(client, import_db):
    headers = await _auth_headers(client)
    response = await _upload(client, headers, "paper.docx", _sample_docx(), DOCX_MIME)
    assert response.status_code == 200, response.text
    payload = response.json()

    assert payload["name"] == "Imported Paper"
    assert payload["mainFile"] == "main.tex"
    assert payload["sourceFormat"] == "docx"
    assert payload["compiler"] == "auto"
    assert payload["warnings"] == []
    assert [f["path"] for f in payload["files"]] == ["main.tex"]
    main = payload["files"][0]["content"]
    assert r"\section{Introduction}" in main
    assert r"Body with 10\% \& more." in main
    assert r"\includegraphics[width=0.40\linewidth]{figures/image1.png}" in main
    assert payload["assets"] == [
        {"name": "figures/image1.png", "mimeType": "image/png", "dataUrl": payload["assets"][0]["dataUrl"]}
    ]
    assert payload["assets"][0]["dataUrl"].startswith("data:image/png;base64,")

    # The payload feeds straight into paper creation (same path as the Overleaf ZIP import).
    created = await client.post(
        "/api/v1/papers",
        headers=headers,
        json={
            "name": payload["name"],
            "latex": main,
            "metadata": {"files": payload["files"], "mainFile": payload["mainFile"], "compiler": payload["compiler"]},
        },
    )
    assert created.status_code == 201, created.text
    patched = await client.patch(
        f"/api/v1/papers/{created.json()['id']}",
        headers=headers,
        json={"assets": payload["assets"]},
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["metadata"]["assets"][0]["name"] == "figures/image1.png"


@pytest.mark.asyncio
async def test_import_pdf_returns_text_only_warning(client, import_db):
    headers = await _auth_headers(client)
    pdf = build_pdf(
        [
            [
                PdfText("Conversion Study", y=80, size=20),
                PdfText("1 Introduction", y=140, size=12, bold=True),
                PdfText("Plain body text.", y=160),
            ]
        ]
    )
    response = await _upload(client, headers, "study.pdf", pdf, "application/pdf")
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["sourceFormat"] == "pdf"
    assert payload["name"] == "Conversion Study"
    assert payload["assets"] == []
    assert [w["code"] for w in payload["warnings"]] == ["pdf_text_only"]
    assert r"\section{Introduction}" in payload["files"][0]["content"]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("filename", "data", "status", "code"),
    [
        ("notes.txt", b"hello", 415, "unsupported_type"),
        ("legacy.doc", b"\xd0\xcf\x11\xe0", 415, "legacy_doc"),
        ("fake.pdf", b"PK\x03\x04 definitely not a pdf", 400, "content_mismatch"),
        ("fake.docx", b"%PDF-1.4 not a word file", 400, "content_mismatch"),
        ("empty.pdf", b"", 400, "empty_file"),
    ],
)
async def test_import_rejects_invalid_files(client, import_db, filename, data, status, code):
    headers = await _auth_headers(client)
    response = await _upload(client, headers, filename, data, "application/octet-stream")
    assert response.status_code == status, response.text
    assert response.json()["detail"]["code"] == code
    assert response.json()["detail"]["message"]


@pytest.mark.asyncio
async def test_import_rejects_files_over_size_limit(client, import_db, monkeypatch):
    monkeypatch.setattr("src.api.import_routes.MAX_IMPORT_BYTES", 1024)
    headers = await _auth_headers(client)
    pdf = b"%PDF-1.4\n" + b"0" * 4096
    response = await _upload(client, headers, "big.pdf", pdf, "application/pdf")
    assert response.status_code == 413
    assert response.json()["detail"]["code"] == "file_too_large"


@pytest.mark.asyncio
async def test_import_reports_unconvertible_pdf(client, import_db):
    headers = await _auth_headers(client)
    response = await _upload(client, headers, "scan.pdf", build_pdf([[]]), "application/pdf")
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "pdf_no_text"
