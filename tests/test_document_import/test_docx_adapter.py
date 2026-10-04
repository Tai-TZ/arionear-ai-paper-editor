import base64
import io
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches

from src.services.document_import import convert_document
from src.services.document_import.docx_adapter import convert_docx
from tests.test_document_import.builders import docx_bytes, image_bytes


def _convert(document, filename: str = "paper.docx"):
    return convert_docx(docx_bytes(document), filename)


def _body(tex: str) -> str:
    return tex.split(r"\maketitle", 1)[1].split(r"\end{document}", 1)[0]


def _add_numbering(paragraph, num_id: int, ilvl: int = 0) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), str(ilvl))
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.append(level)
    num_pr.append(number)
    ppr.append(num_pr)


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def test_headings_title_and_sections():
    doc = Document()
    doc.add_paragraph("Graph Networks for 50% Speedups", style="Title")
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("Intro text.")
    doc.add_heading("Methods", level=2)
    doc.add_heading("Datasets", level=3)
    doc.add_heading("Fine print", level=4)

    project = _convert(doc)
    tex = project.main_content

    assert project.name == "Graph Networks for 50% Speedups"
    assert project.main_file == "main.tex"
    assert project.files[0].path == "main.tex"
    assert project.source_format == "docx"
    assert project.compiler == "auto"
    assert r"\title{Graph Networks for 50\% Speedups}" in tex
    assert r"\section{Introduction}" in tex  # manual numbering stripped
    assert r"\subsection{Methods}" in tex
    assert r"\subsubsection{Datasets}" in tex
    assert r"\paragraph{Fine print}" in tex
    assert tex.index(r"\section{Introduction}") < tex.index("Intro text.") < tex.index(r"\subsection{Methods}")


def test_first_heading_becomes_title_when_no_title_style():
    doc = Document()
    doc.add_heading("Robust Parsing of Documents", level=1)
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Body.")
    project = _convert(doc)
    assert project.name == "Robust Parsing of Documents"
    assert r"\title{Robust Parsing of Documents}" in project.main_content
    assert r"\section{Robust Parsing" not in project.main_content
    assert r"\section{Introduction}" in project.main_content


def test_filename_is_fallback_title():
    doc = Document()
    doc.add_paragraph("Only a paragraph.")
    project = _convert(doc, "my_draft_v2.docx")
    assert project.name == "my draft v2"
    assert r"\title{my draft v2}" in project.main_content


def test_inline_formatting_is_merged_and_escaped():
    doc = Document()
    paragraph = doc.add_paragraph("Plain ")
    paragraph.add_run("bold ").bold = True
    paragraph.add_run("still bold").bold = True
    paragraph.add_run(", ")
    paragraph.add_run("italic").italic = True
    paragraph.add_run(" and ")
    paragraph.add_run("under").underline = True
    paragraph.add_run(" H")
    paragraph.add_run("2").font.subscript = True
    paragraph.add_run("O x")
    paragraph.add_run("2").font.superscript = True
    both = paragraph.add_run(" 100% & $")
    both.bold = True
    both.italic = True

    tex = _convert(doc).main_content
    assert r"Plain \textbf{bold still bold}, \textit{italic} and \underline{under}" in tex
    assert r"H\textsubscript{2}O x\textsuperscript{2}" in tex
    assert r"\textbf{\textit{100\% \& \$}}" in tex


def test_strong_and_emphasis_character_styles():
    doc = Document()
    paragraph = doc.add_paragraph("A ")
    paragraph.add_run("key", style="Strong")
    paragraph.add_run(" and ")
    paragraph.add_run("nuance", style="Emphasis")
    tex = _convert(doc).main_content
    assert r"A \textbf{key} and \textit{nuance}" in tex


def test_line_breaks_inside_paragraph_use_newline():
    doc = Document()
    paragraph = doc.add_paragraph("first")
    paragraph.add_run().add_break()
    paragraph.add_run("second")
    tex = _convert(doc).main_content
    assert r"first\newline second" in tex


def test_bulleted_numbered_and_nested_lists():
    doc = Document()
    doc.add_paragraph("Before list.")
    doc.add_paragraph("apple", style="List Bullet")
    doc.add_paragraph("banana", style="List Bullet")
    doc.add_paragraph("banana split", style="List Bullet 2")
    doc.add_paragraph("one", style="List Number")
    doc.add_paragraph("two", style="List Number")
    doc.add_paragraph("After list.")

    body = _body(_convert(doc).main_content)
    expected = "\n".join(
        [
            r"\begin{itemize}",
            r"\item apple",
            r"\item banana",
            r"\begin{itemize}",
            r"\item banana split",
            r"\end{itemize}",
            r"\end{itemize}",
            r"\begin{enumerate}",
            r"\item one",
            r"\item two",
            r"\end{enumerate}",
        ]
    )
    assert expected in body
    assert body.index("Before list.") < body.index(r"\begin{itemize}")
    assert body.index(r"\end{enumerate}") < body.index("After list.")


def test_lists_detected_from_paragraph_numbering_properties():
    doc = Document()
    bullet_num_id = doc.styles["List Bullet"].element.pPr.numPr.numId.val
    decimal_num_id = doc.styles["List Number"].element.pPr.numPr.numId.val
    _add_numbering(doc.add_paragraph("dot item"), bullet_num_id)
    _add_numbering(doc.add_paragraph("counted item"), decimal_num_id)
    _add_numbering(doc.add_paragraph("counted child"), decimal_num_id, ilvl=1)

    body = _body(_convert(doc).main_content)
    assert "\\begin{itemize}\n\\item dot item\n\\end{itemize}" in body
    assert "\\begin{enumerate}\n\\item counted item\n\\begin{enumerate}\n\\item counted child" in body
    assert body.count(r"\begin{enumerate}") == body.count(r"\end{enumerate}") == 2


def test_table_becomes_booktabs_tabular_with_caption_and_merged_cells():
    doc = Document()
    doc.add_paragraph("Table 1: Accuracy by model", style="Caption")
    table = doc.add_table(rows=3, cols=3)
    for col, text in enumerate(["Model", "Acc (%)", "F_1"]):
        table.cell(0, col).text = text
    for col, text in enumerate(["Base", "91.2", "0.88"]):
        table.cell(1, col).text = text
    merged = table.cell(2, 0).merge(table.cell(2, 1))
    merged.text = "Average & more"
    table.cell(2, 2).text = "0.9"
    table.cell(1, 0).paragraphs[0].runs[0].bold = True

    tex = _convert(doc).main_content
    assert "\\begin{table}[htbp]\n\\centering\n\\caption{Accuracy by model}\n\\begin{tabular}{lll}" in tex
    assert r"Model & Acc (\%) & F\_1 \\" in tex
    assert r"\textbf{Base} & 91.2 & 0.88 \\" in tex
    assert r"\multicolumn{2}{l}{Average \& more} & 0.9 \\" in tex
    assert tex.index(r"\toprule") < tex.index(r"\midrule") < tex.index(r"\bottomrule")
    assert "Table 1:" not in tex


def test_vertically_merged_cells_are_not_repeated():
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "H1"
    table.cell(0, 1).text = "H2"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "tall"
    table.cell(1, 1).text = "a"
    table.cell(2, 1).text = "b"
    tex = _convert(doc).main_content
    assert "tall & a \\\\\n & b \\\\" in tex


def test_inline_images_are_extracted_as_figure_assets():
    doc = Document()
    doc.add_paragraph("See the figure.")
    doc.add_picture(io.BytesIO(image_bytes("PNG")), width=Inches(2.5))
    doc.add_paragraph("Figure 1: A red box", style="Caption")
    doc.add_picture(io.BytesIO(image_bytes("GIF", color="blue")), width=Inches(8))

    project = _convert(doc)
    tex = project.main_content
    names = [asset.name for asset in project.assets]
    assert names == ["figures/image1.png", "figures/image2.png"]  # GIF converted to PNG
    assert all(asset.mime_type == "image/png" for asset in project.assets)
    assert project.assets[0].data.startswith(b"\x89PNG")
    assert project.assets[1].data.startswith(b"\x89PNG")
    data_url = project.assets[0].data_url
    assert data_url.startswith("data:image/png;base64,")
    assert base64.b64decode(data_url.split(",", 1)[1]) == project.assets[0].data

    assert r"\includegraphics[width=0.50\linewidth]{figures/image1.png}" in tex
    assert r"\includegraphics[width=\linewidth]{figures/image2.png}" in tex
    figure = tex[tex.index(r"\begin{figure}") : tex.index(r"\end{figure}")]
    assert r"\caption{A red box}" in figure
    assert project.warnings == []


def test_same_image_used_twice_is_stored_once():
    doc = Document()
    png = image_bytes("PNG")
    doc.add_picture(io.BytesIO(png))
    doc.add_picture(io.BytesIO(png))
    project = _convert(doc)
    assert len(project.assets) == 1
    assert project.main_content.count(r"\includegraphics") == 2


def test_abstract_heading_becomes_abstract_environment():
    doc = Document()
    doc.add_paragraph("Paper", style="Title")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("We propose a method.")
    doc.add_paragraph("It works.")
    doc.add_heading("Introduction", level=1)
    tex = _convert(doc).main_content
    assert "\\begin{abstract}\nWe propose a method.\n\nIt works.\n\\end{abstract}" in tex
    assert r"\section{Abstract}" not in tex
    assert tex.index(r"\end{abstract}") < tex.index(r"\section{Introduction}")


def test_abstract_paragraph_with_inline_label():
    doc = Document()
    doc.add_paragraph("Abstract: Short summary_here.")
    doc.add_paragraph("Body text.")
    tex = _convert(doc).main_content
    assert "\\begin{abstract}\nShort summary\\_here.\n\\end{abstract}\n\nBody text." in tex


def test_hyperlinks_become_href():
    doc = Document()
    paragraph = doc.add_paragraph("Visit ")
    _add_hyperlink(paragraph, "https://example.org/a_b?x=1%20#top", "our site")
    tex = _convert(doc).main_content
    assert r"Visit \href{https://example.org/a_b?x=1\%20\#top}{our site}" in tex


def test_equations_are_flagged_with_placeholder_and_warning():
    doc = Document()
    paragraph = doc.add_paragraph("Energy is ")
    paragraph._p.append(
        parse_xml(
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:r><m:t>E=mc^2</m:t></m:r></m:oMath>"
        )
    )
    project = _convert(doc)
    assert r"Energy is \textbf{[equation omitted]}" in project.main_content
    codes = [warning.code for warning in project.warnings]
    assert codes == ["docx_equations_skipped"]
    assert project.warnings[0].count == 1
    assert "% - 1 Word equation(s) could not be converted" in project.main_content


def test_vietnamese_text_switches_to_unicode_engine():
    doc = Document()
    doc.add_paragraph("Phương pháp nghiên cứu", style="Title")
    doc.add_paragraph("Kết quả cho thấy độ chính xác cao.")
    project = _convert(doc)
    assert project.compiler == "xelatex"
    assert r"\usepackage{fontspec}" in project.main_content
    assert r"\usepackage[T1]{fontenc}" not in project.main_content
    assert project.name == "Phương pháp nghiên cứu"


def test_empty_document_warns():
    project = _convert(Document())
    assert [warning.code for warning in project.warnings] == ["empty_document"]
    assert r"\begin{document}" in project.main_content


def test_convert_document_dispatches_to_docx():
    doc = Document()
    doc.add_paragraph("Dispatch check.")
    project = convert_document(docx_bytes(doc), "dispatch.docx")
    assert "Dispatch check." in project.main_content


@pytest.mark.skipif(shutil.which("pdflatex") is None, reason="pdflatex not installed")
def test_generated_latex_compiles_with_pdflatex(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Compile Check: 100% & More", style="Title")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("Special chars \\ { } ~ ^ _ # $ and α ≤ β.")
    doc.add_heading("Introduction", level=1)
    paragraph = doc.add_paragraph("Mixed ")
    paragraph.add_run("bold").bold = True
    paragraph.add_run().add_break()
    paragraph.add_run("[1] after a break")
    doc.add_paragraph("[1] item", style="List Bullet")
    doc.add_paragraph("nested", style="List Bullet 3")
    doc.add_paragraph("number", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A_1"
    table.cell(0, 1).text = "B & C"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "merged"
    doc.add_picture(io.BytesIO(image_bytes("PNG")), width=Inches(2))
    doc.add_paragraph("Figure 1: Caption", style="Caption")

    project = _convert(doc)
    (tmp_path / "main.tex").write_text(project.main_content, encoding="utf-8")
    for asset in project.assets:
        target = tmp_path / asset.name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(asset.data)
    try:
        result = subprocess.run(
            ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "main.tex"],
            cwd=tmp_path,
            capture_output=True,
            timeout=180,
        )
    except subprocess.TimeoutExpired:
        pytest.skip("pdflatex timed out")
    log = result.stdout.decode("utf-8", errors="replace")
    if re.search(r"File `[^']+\.sty' not found", log):
        pytest.skip("TeX installation lacks a required package")
    assert result.returncode == 0, log[-3000:]
    assert (tmp_path / "main.pdf").exists()
